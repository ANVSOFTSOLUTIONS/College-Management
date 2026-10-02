"""NAAC / AISHE data: the numbers accreditation forms ask for, from data the ERP already holds.

One report of tables, each tagged with the NAAC criterion it feeds, for the
college's current (non-archived) batches. Downloadable as an Excel workbook
with a sheet per table.
"""

import io
import re
from datetime import date
from statistics import median

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel

from app.api.deps import CurrentUser
from app.db.helpers import fetch_all, fetch_one
from app.modules.exams.service import class_results
from app.modules.grievances.service import CATEGORY_LABELS

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
Cell = str | int | float | None


class Metric(BaseModel):
    label: str
    value: Cell


class Table(BaseModel):
    key: str
    title: str
    criterion: str
    headers: list[str]
    rows: list[list[Cell]]


class NaacReport(BaseModel):
    college_name: str
    generated_on: date
    metrics: list[Metric]
    tables: list[Table]


def _pct(part: int | float, whole: int | float) -> float | None:
    return round(part * 100 / whole, 1) if whole else None


def _designation_group(designation: str) -> str:
    d = designation.lower()
    if "assistant" in d:
        return "assistant"
    if "associate" in d:
        return "associate"
    if "professor" in d:
        return "professor"
    return "other"


async def _students(school_id: str) -> tuple[list[Table], int]:
    rows = await fetch_all(
        """
        SELECT s.gender, s.quota, s.social_category, c.id AS class_id, c.program, COALESCE(d.name, 'No department') AS department
        FROM students s JOIN classes c ON c.id = s.class_id LEFT JOIN departments d ON d.id = c.department_id
        WHERE s.school_id = %s AND s.status = 'active' AND c.is_archived = 0
        """,
        (school_id,),
    )
    by_program: dict[tuple, dict] = {}
    for r in rows:
        g = by_program.setdefault((r["department"], r["program"] or "-"), {"batches": set(), "male": 0, "female": 0, "other": 0, "total": 0})
        g["batches"].add(r["class_id"])
        g[r["gender"] if r["gender"] in ("male", "female") else "other"] += 1
        g["total"] += 1
    programs = Table(
        key="students_programmes", title="Students by department and programme", criterion="2.1 / AISHE",
        headers=["Department", "Programme", "Batches", "Male", "Female", "Other / not set", "Total"],
        rows=[[d, p, len(g["batches"]), g["male"], g["female"], g["other"], g["total"]] for (d, p), g in sorted(by_program.items())],
    )
    by_category: dict[str, list[int]] = {}
    for r in rows:
        c = by_category.setdefault(r["social_category"] or "Not set", [0, 0, 0])
        c[{"male": 0, "female": 1}.get(r["gender"], 2)] += 1
    categories = Table(
        key="students_categories", title="Students by social category", criterion="2.1.2 / AISHE",
        headers=["Category", "Male", "Female", "Other / not set", "Total"],
        rows=[[c, *v, sum(v)] for c, v in sorted(by_category.items())],
    )
    by_quota: dict[str, int] = {}
    for r in rows:
        by_quota[r["quota"] or "Not set"] = by_quota.get(r["quota"] or "Not set", 0) + 1
    quotas = Table(key="students_quota", title="Students by admission quota", criterion="2.1", headers=["Quota", "Students"],
                   rows=[[q, n] for q, n in sorted(by_quota.items())])
    return [programs, categories, quotas], len(rows)


async def _faculty(school_id: str) -> tuple[Table, int]:
    rows = await fetch_all(
        """
        SELECT t.designation, COALESCE(d.name, NULLIF(t.department, ''), 'No department') AS department
        FROM teachers t JOIN users u ON u.id = t.user_id LEFT JOIN departments d ON d.id = t.department_id
        WHERE t.school_id = %s AND u.status = 'active'
        """,
        (school_id,),
    )
    groups: dict[str, dict] = {}
    for r in rows:
        g = groups.setdefault(r["department"], {"professor": 0, "associate": 0, "assistant": 0, "other": 0})
        g[_designation_group(r["designation"] or "")] += 1
    table = Table(
        key="faculty", title="Full-time faculty by department and designation", criterion="2.4",
        headers=["Department", "Professors", "Associate Professors", "Assistant Professors", "Other", "Total"],
        rows=[[d, g["professor"], g["associate"], g["assistant"], g["other"], sum(g.values())] for d, g in sorted(groups.items())],
    )
    return table, len(rows)


async def _results(user: CurrentUser) -> tuple[Table, int, int]:
    """Each current batch's latest published semester-end exam."""
    rows = await fetch_all(
        """
        SELECT c.id AS class_id, c.name, c.section, c.semester,
               (SELECT e.id FROM exams e JOIN exam_subjects es ON es.exam_id = e.id
                WHERE es.class_id = c.id AND e.exam_type = 'semester' AND e.published_at IS NOT NULL
                ORDER BY e.published_at DESC LIMIT 1) AS exam_id
        FROM classes c WHERE c.school_id = %s AND c.is_archived = 0 ORDER BY c.name, c.section
        """,
        (user.school_id,),
    )
    table_rows, appeared_all, passed_all = [], 0, 0
    for r in rows:
        if not r["exam_id"]:
            continue
        results = await class_results(user, r["exam_id"], r["class_id"])
        complete = [s for s in results.students if s.complete]
        passed = sum(1 for s in complete if s.passed)
        appeared_all += len(complete)
        passed_all += passed
        table_rows.append([f"{r['name']} - {r['section']}", r["semester"], results.exam_name, len(complete), passed, _pct(passed, len(complete))])
    table = Table(key="results", title="Pass percentage (latest semester-end exam per batch)", criterion="2.6.3",
                  headers=["Batch", "Semester", "Exam", "Appeared", "Passed", "Pass %"], rows=table_rows)
    return table, appeared_all, passed_all


async def _placements(school_id: str) -> tuple[Table, int, float | None, float | None]:
    rows = await fetch_all(
        """
        SELECT pc.name AS company, d.role_title, d.package_lpa, d.drive_date,
               COUNT(a.id) AS applied, SUM(a.status = 'selected') AS selected
        FROM placement_drives d JOIN placement_companies pc ON pc.id = d.company_id
        LEFT JOIN placement_applications a ON a.drive_id = d.id
        WHERE d.school_id = %s GROUP BY d.id, pc.name, d.role_title, d.package_lpa, d.drive_date ORDER BY d.drive_date DESC, pc.name
        """,
        (school_id,),
    )
    placed = await fetch_all(
        """
        SELECT a.student_id, MAX(d.package_lpa) AS package FROM placement_applications a JOIN placement_drives d ON d.id = a.drive_id
        WHERE a.school_id = %s AND a.status = 'selected' GROUP BY a.student_id
        """,
        (school_id,),
    )
    packages = [float(p["package"]) for p in placed if p["package"] is not None]
    table = Table(
        key="placements", title="Placement drives", criterion="5.2.1",
        headers=["Company", "Role", "Drive date", "Package (LPA)", "Applied", "Selected"],
        rows=[[r["company"], r["role_title"], r["drive_date"].isoformat() if r["drive_date"] else None,
               float(r["package_lpa"]) if r["package_lpa"] is not None else None, r["applied"], int(r["selected"] or 0)] for r in rows],
    )
    return table, len(placed), max(packages) if packages else None, round(median(packages), 2) if packages else None


async def _feedback(school_id: str) -> tuple[Table, float | None]:
    latest = await fetch_one(
        "SELECT id, title FROM feedback_rounds WHERE school_id = %s ORDER BY is_open, created_at DESC LIMIT 1", (school_id,)
    )
    if latest is None:
        return Table(key="feedback", title="Student feedback on faculty", criterion="1.4", headers=["Department", "Faculty rated", "Responses", "Average (out of 5)"], rows=[]), None
    rows = await fetch_all(
        """
        SELECT COALESCE(d.name, NULLIF(t.department, ''), 'No department') AS department, f.teacher_id, f.ratings
        FROM feedback_responses f JOIN teachers t ON t.id = f.teacher_id LEFT JOIN departments d ON d.id = t.department_id
        WHERE f.round_id = %s
        """,
        (latest["id"],),
    )
    groups: dict[str, dict] = {}
    everything = []
    for r in rows:
        score = sum(int(x) for x in r["ratings"].split(",")) / len(r["ratings"].split(","))
        everything.append(score)
        g = groups.setdefault(r["department"], {"faculty": set(), "scores": []})
        g["faculty"].add(r["teacher_id"])
        g["scores"].append(score)
    table = Table(
        key="feedback", title=f"Student feedback on faculty ({latest['title']})", criterion="1.4",
        headers=["Department", "Faculty rated", "Responses", "Average (out of 5)"],
        rows=[[d, len(g["faculty"]), len(g["scores"]), round(sum(g["scores"]) / len(g["scores"]), 2)] for d, g in sorted(groups.items())],
    )
    return table, round(sum(everything) / len(everything), 2) if everything else None


async def _grievances(school_id: str) -> tuple[Table, int, int]:
    rows = await fetch_all(
        """
        SELECT category, COUNT(*) AS received, SUM(status IN ('resolved', 'closed')) AS resolved,
               AVG(CASE WHEN resolved_at IS NOT NULL THEN TIMESTAMPDIFF(HOUR, created_at, resolved_at) END) AS hours
        FROM grievances WHERE school_id = %s GROUP BY category ORDER BY category
        """,
        (school_id,),
    )
    table = Table(
        key="grievances", title="Grievance redressal", criterion="5.1.5",
        headers=["Category", "Received", "Resolved", "Pending", "Average days to resolve"],
        rows=[[CATEGORY_LABELS.get(r["category"], r["category"]), r["received"], int(r["resolved"] or 0), r["received"] - int(r["resolved"] or 0),
               round(float(r["hours"]) / 24, 1) if r["hours"] is not None else None] for r in rows],
    )
    return table, sum(r["received"] for r in rows), sum(int(r["resolved"] or 0) for r in rows)


async def _attendance(school_id: str) -> Table:
    rows = await fetch_all(
        """
        SELECT c.name, c.section, a.student_id, COUNT(*) AS held, SUM(a.status IN ('present', 'late')) AS attended
        FROM subject_attendance a JOIN classes c ON c.id = a.class_id
        WHERE a.school_id = %s AND c.is_archived = 0 GROUP BY c.id, c.name, c.section, a.student_id
        """,
        (school_id,),
    )
    groups: dict[str, list[float]] = {}
    for r in rows:
        groups.setdefault(f"{r['name']} - {r['section']}", []).append(int(r["attended"]) * 100 / r["held"])
    return Table(
        key="attendance", title="Student attendance (subject-wise)", criterion="2.3",
        headers=["Batch", "Students", "Average attendance %", "Students below 75%"],
        rows=[[b, len(v), round(sum(v) / len(v), 1), sum(1 for x in v if x < 75)] for b, v in sorted(groups.items())],
    )


async def _alumni(school_id: str) -> Table:
    rows = await fetch_all(
        """
        SELECT passing_year, COUNT(*) AS total, SUM(status = 'employed') AS employed, SUM(status = 'higher_studies') AS higher,
               SUM(status = 'self_employed') AS self_employed
        FROM alumni WHERE school_id = %s GROUP BY passing_year ORDER BY passing_year DESC
        """,
        (school_id,),
    )
    return Table(
        key="alumni", title="Alumni progression", criterion="5.2.2 / 5.4",
        headers=["Passing year", "Alumni", "Employed", "Higher studies", "Self-employed"],
        rows=[[r["passing_year"], r["total"], int(r["employed"] or 0), int(r["higher"] or 0), int(r["self_employed"] or 0)] for r in rows],
    )


async def report(user: CurrentUser) -> NaacReport:
    school = await fetch_one("SELECT name FROM schools WHERE id = %s", (user.school_id,))
    student_tables, students = await _students(user.school_id)
    faculty_table, faculty = await _faculty(user.school_id)
    results_table, appeared, passed = await _results(user)
    placement_table, placed, highest, median_package = await _placements(user.school_id)
    feedback_table, feedback_avg = await _feedback(user.school_id)
    grievance_table, received, resolved = await _grievances(user.school_id)
    departments = await fetch_one("SELECT COUNT(*) AS n FROM departments WHERE school_id = %s", (user.school_id,))
    metrics = [
        Metric(label="Students (current batches)", value=students),
        Metric(label="Full-time faculty", value=faculty),
        Metric(label="Student : faculty ratio", value=f"{round(students / faculty, 1)} : 1" if faculty else None),
        Metric(label="Departments", value=departments["n"]),
        Metric(label="Pass percentage", value=_pct(passed, appeared)),
        Metric(label="Students placed", value=placed),
        Metric(label="Highest package (LPA)", value=highest),
        Metric(label="Median package (LPA)", value=median_package),
        Metric(label="Average faculty feedback (out of 5)", value=feedback_avg),
        Metric(label="Grievances resolved", value=f"{resolved} of {received}"),
    ]
    tables = [*student_tables, faculty_table, results_table, await _attendance(user.school_id), placement_table, await _alumni(user.school_id), feedback_table, grievance_table]
    return NaacReport(college_name=school["name"], generated_on=date.today(), metrics=metrics, tables=tables)


def to_xlsx(data: NaacReport) -> StreamingResponse:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append([data.college_name])
    summary.append([f"NAAC / AISHE data as of {data.generated_on.isoformat()}"])
    summary.append([])
    for m in data.metrics:
        summary.append([m.label, m.value])
    summary["A1"].font = Font(bold=True, size=14)
    summary.column_dimensions["A"].width = 40
    summary.column_dimensions["B"].width = 18
    for table in data.tables:
        sheet = workbook.create_sheet(re.sub(r"[\[\]:*?/\\]", "", table.title)[:31])
        sheet.append([f"{table.title} (Criterion {table.criterion})"])
        sheet["A1"].font = Font(bold=True)
        sheet.append(table.headers)
        for cell in sheet[2]:
            cell.font = Font(bold=True)
        for row in table.rows:
            sheet.append(row)
        for column in sheet.iter_cols(min_row=2):
            width = max(len(str(c.value)) if c.value is not None else 0 for c in column)
            sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 45)
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type=XLSX, headers={"Content-Disposition": 'attachment; filename="naac-aishe-data.xlsx"', "Cache-Control": "no-store"})
