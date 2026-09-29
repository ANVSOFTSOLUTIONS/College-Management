"""Bulk import of students (with parents) from an Excel or CSV file.

The admin downloads a template, fills it in, and uploads it. Every row is
validated exactly like a student added by hand; a dry run returns the row-by-
row result without saving. The import itself saves all valid rows in one
transaction, or refuses if any row is invalid unless told to skip those.
"""

import csv
import io
import re
from datetime import date, datetime

import aiomysql
from fastapi import UploadFile, status
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, ValidationError

from app.api.deps import CurrentUser
from app.core.errors import AppError
from app.db.database import db
from app.db.helpers import fetch_all
from app.modules.students.schemas import CreateStudentRequest
from app.modules.students.service import insert_student

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 2000

# (field, header, required, example)
COLUMNS = [
    ("admission_number", "Admission number", True, "2026-001"),
    ("full_name", "Full name", True, "Ananya Rao"),
    ("class_name", "Class", True, "Grade 5"),
    ("section", "Section", True, "A"),
    ("date_of_birth", "Date of birth (DD-MM-YYYY)", False, "12-04-2016"),
    ("gender", "Gender (Male/Female/Other)", False, "Female"),
    ("blood_group", "Blood group", False, "B+"),
    ("admission_date", "Admission date (DD-MM-YYYY)", False, "01-06-2026"),
    ("address", "Address", False, "12-3, Madhapur, Hyderabad"),
    ("father_name", "Father name", False, "Srinivas Rao"),
    ("father_phone", "Father phone", False, "9848022338"),
    ("father_email", "Father email", False, ""),
    ("mother_name", "Mother name", False, "Padma Rao"),
    ("mother_phone", "Mother phone", False, "9848022339"),
    ("mother_email", "Mother email", False, ""),
    ("guardian_name", "Guardian name", False, ""),
    ("guardian_relation", "Guardian relation", False, ""),
    ("guardian_phone", "Guardian phone", False, ""),
    ("primary_contact", "Primary contact (Father/Mother/Guardian)", False, "Father"),
]


class ImportRow(BaseModel):
    row: int
    admission_number: str
    full_name: str
    class_label: str
    errors: list[str]


class ImportResult(BaseModel):
    total: int
    valid: int
    invalid: int
    imported: int
    rows: list[ImportRow]


def _key(header) -> str:
    """'Date of birth (DD-MM-YYYY)*' → 'date of birth'."""
    text = re.sub(r"\(.*?\)", "", str(header or "")).replace("*", "")
    return re.sub(r"\s+", " ", text).strip().lower()


HEADER_TO_FIELD = {_key(header): field for field, header, _, _ in COLUMNS}


# --- Template -------------------------------------------------------------------


async def template_bytes(school_id: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Students"
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="059669")
    for index, (_, header, required, example) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=f"{header}{' *' if required else ''}")
        cell.font, cell.fill = bold, fill
        sheet.cell(row=2, column=index, value=example)
        sheet.column_dimensions[cell.column_letter].width = max(14, len(header) + 4)
    sheet.freeze_panes = "A2"

    classes = await fetch_all(
        "SELECT name, section, academic_year FROM classes WHERE school_id = %s ORDER BY academic_year DESC, name, section", (school_id,)
    )
    info = workbook.create_sheet("Classes")
    info.append(["Use these exact Class and Section values"])
    info.append(["Class", "Section", "Academic year"])
    for row in classes:
        info.append([row["name"], row["section"], row["academic_year"]])
    info.column_dimensions["A"].width = 30
    notes = workbook.create_sheet("How to fill")
    for line in [
        "Row 2 of the Students sheet is an example: replace or delete it.",
        "Columns marked * are required.",
        "Phone numbers: 10-digit mobile numbers. Give at least one parent's phone so alerts and the parent app work.",
        "Dates: DD-MM-YYYY, e.g. 12-04-2016.",
        "Primary contact: who gets alerts (Father, Mother, or Guardian). Leave blank to use the father, else the mother.",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 110

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- Reading ----------------------------------------------------------------------


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # phone numbers typed into Excel become floats
    if isinstance(value, (datetime, date)):
        return value.strftime("%d-%m-%Y")
    return str(value).strip()


def _read_rows(filename: str, content: bytes) -> list[dict]:
    if filename.lower().endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        raw = list(csv.reader(io.StringIO(text)))
    else:
        try:
            sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
        except Exception as exc:  # corrupt or not an Excel file
            raise AppError(status.HTTP_400_BAD_REQUEST, "invalid_file", "Couldn't read this file. Upload the Excel (.xlsx) template or a CSV.") from exc
        raw = [list(row) for row in sheet.iter_rows(values_only=True)]
    if not raw:
        raise AppError(status.HTTP_400_BAD_REQUEST, "empty_file", "The file is empty.")

    fields = [HEADER_TO_FIELD.get(_key(h)) for h in raw[0]]
    missing = [header for field, header, required, _ in COLUMNS if required and field not in fields]
    if missing:
        raise AppError(status.HTTP_400_BAD_REQUEST, "missing_columns", f"Missing column(s): {', '.join(missing)}. Use the template.")

    rows = []
    for number, values in enumerate(raw[1:], start=2):
        record = {field: _cell_text(value) for field, value in zip(fields, values) if field}
        if any(record.values()):
            record["_row"] = number
            rows.append(record)
    if len(rows) > MAX_ROWS:
        raise AppError(status.HTTP_400_BAD_REQUEST, "too_many_rows", f"Import at most {MAX_ROWS} students at a time.")
    return rows


def _parse_date(text: str) -> date | None:
    if not text:
        return None
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"'{text}' isn't a date in DD-MM-YYYY format")


GENDERS = {"male": "male", "m": "male", "boy": "male", "female": "female", "f": "female", "girl": "female", "other": "other"}


def _guardian(record: dict, prefix: str) -> dict | None:
    name = record.get(f"{prefix}_name", "")
    if not name:
        if record.get(f"{prefix}_phone") or record.get(f"{prefix}_email"):
            raise ValueError(f"{prefix.capitalize()} name is missing")
        return None
    guardian = {"full_name": name, "phone": record.get(f"{prefix}_phone", ""), "email": record.get(f"{prefix}_email", "")}
    if prefix == "guardian":
        guardian["relation_label"] = record.get("guardian_relation", "")
    return guardian


def _validate(record: dict, class_ids: dict, existing: set, seen: set) -> tuple[CreateStudentRequest | None, list[str]]:
    errors = []
    class_label = f"{record.get('class_name', '')} - {record.get('section', '')}"
    class_id = class_ids.get((record.get("class_name", "").lower(), record.get("section", "").lower()))
    if record.get("class_name") and record.get("section") and class_id is None:
        errors.append(f"Class '{class_label}' not found; add it under Classes & Subjects first")
    admission = record.get("admission_number", "").strip()
    if admission.lower() in existing:
        errors.append(f"Admission number {admission} already exists")
    elif admission.lower() in seen:
        errors.append(f"Admission number {admission} appears twice in the file")

    payload = None
    try:
        gender = record.get("gender", "").lower()
        if gender and gender not in GENDERS:
            raise ValueError(f"Gender '{record['gender']}' should be Male, Female or Other")
        primary = record.get("primary_contact", "").lower()
        if primary and primary not in ("father", "mother", "guardian"):
            raise ValueError("Primary contact should be Father, Mother or Guardian")
        payload = CreateStudentRequest(
            admission_number=admission,
            full_name=record.get("full_name", ""),
            class_id=class_id or "unknown",
            date_of_birth=_parse_date(record.get("date_of_birth", "")),
            gender=GENDERS.get(gender, ""),
            blood_group=record.get("blood_group", ""),
            admission_date=_parse_date(record.get("admission_date", "")),
            address=record.get("address", ""),
            father=_guardian(record, "father"),
            mother=_guardian(record, "mother"),
            guardian=_guardian(record, "guardian"),
            primary_contact=primary,
        )
    except ValueError as exc:  # includes pydantic ValidationError
        if isinstance(exc, ValidationError):
            for err in exc.errors():
                where = ".".join(str(p) for p in err["loc"]).replace("_", " ")
                errors.append(f"{where}: {err['msg'].removeprefix('Value error, ')}" if where else err["msg"].removeprefix("Value error, "))
        else:
            errors.append(str(exc))
    return (payload if not errors else None), errors


async def run_import(user: CurrentUser, file: UploadFile, *, dry_run: bool, skip_invalid: bool) -> ImportResult:
    content = await file.read(MAX_BYTES + 1)
    if len(content) > MAX_BYTES:
        raise AppError(status.HTTP_400_BAD_REQUEST, "file_too_large", "Files must be 2MB or smaller.")
    records = _read_rows(file.filename or "", content)

    classes = await fetch_all(
        "SELECT id, name, section FROM classes WHERE school_id = %s ORDER BY academic_year", (user.school_id,)
    )
    class_ids = {(c["name"].lower(), c["section"].lower()): c["id"] for c in classes}  # latest year wins
    existing = {r["n"].lower() for r in await fetch_all("SELECT admission_number AS n FROM students WHERE school_id = %s", (user.school_id,))}

    seen: set[str] = set()
    rows, valid = [], []
    for record in records:
        payload, errors = _validate(record, class_ids, existing, seen)
        seen.add(record.get("admission_number", "").strip().lower())
        rows.append(
            ImportRow(
                row=record["_row"],
                admission_number=record.get("admission_number", ""),
                full_name=record.get("full_name", ""),
                class_label=f"{record.get('class_name', '')} - {record.get('section', '')}",
                errors=errors,
            )
        )
        if payload:
            valid.append(payload)

    result = ImportResult(total=len(rows), valid=len(valid), invalid=len(rows) - len(valid), imported=0, rows=rows)
    if dry_run or not valid:
        return result
    if result.invalid and not skip_invalid:
        raise AppError(status.HTTP_400_BAD_REQUEST, "rows_invalid", f"{result.invalid} row(s) have problems. Fix them, or import only the valid rows.")

    async with db.pool.acquire() as conn:
        await conn.begin()
        try:
            async with conn.cursor() as cur:
                for payload in valid:
                    await insert_student(cur, user.school_id, payload)
        except aiomysql.IntegrityError as exc:
            await conn.rollback()
            raise AppError(status.HTTP_409_CONFLICT, "import_conflict", "Some admission numbers were added meanwhile. Check the file again.") from exc
        await conn.commit()
    result.imported = len(valid)
    return result
