"""Grievance tickets from the app, and the NAAC / AISHE data report."""

from openpyxl import load_workbook
import io

from tests.factories import auth_headers, create_school, create_user, login

PASSWORD = "Secret123!"
API = "/api/v1"


async def _college(client, code):
    school_id = await create_school(code=code)
    await create_user(school_id=school_id, email=f"{code}-admin@example.com".lower(), password=PASSWORD, role="admin")
    admin = auth_headers((await login(client, f"{code}-admin@example.com".lower(), PASSWORD)).json()["access_token"])
    cse = (await client.post(f"{API}/departments", json={"name": "Computer Science", "code": "CSE"}, headers=admin)).json()
    body = {"email": f"{code}-f@example.com".lower(), "full_name": "Dr. Rao", "password": PASSWORD, "department_id": cse["id"], "designation": "Professor"}
    faculty = (await client.post(f"{API}/teachers", json=body, headers=admin)).json()
    batch = (await client.post(f"{API}/classes", json={"name": "B.Tech CSE", "section": "A", "academic_year": "2026", "class_teacher_id": faculty["id"],
                                                        "department_id": cse["id"], "program": "B.Tech", "semester": 1}, headers=admin)).json()
    student = (await client.post(f"{API}/students", json={"admission_number": "S1", "full_name": "Asha", "class_id": batch["id"], "gender": "female",
                                                           "quota": "Convener", "social_category": "BC-A"}, headers=admin)).json()
    await client.post(f"{API}/students/{student['id']}/login", json={"password": "StudentPass1"}, headers=admin)
    token = (await client.post(f"{API}/auth/login", json={"college_code": code, "roll_number": "S1", "password": "StudentPass1"})).json()["access_token"]
    return admin, auth_headers(token), student, batch, faculty


async def test_grievance_ticket_flow(db, client):
    admin, asha_h, asha, batch, faculty = await _college(client, "GRV1")
    created = await client.post(f"{API}/grievances", json={"category": "ragging", "subject": "Seniors in hostel", "description": "Seniors forced us to stand all night."},
                                headers=asha_h)
    assert created.status_code == 201, created.text
    ticket = created.json()
    assert (ticket["ticket_number"], ticket["priority"], ticket["status"], ticket["student_name"]) == ("GRV-00001", "high", "open", "Asha")
    second = (await client.post(f"{API}/grievances", json={"category": "fees", "subject": "Receipt missing", "description": "Paid fees but no receipt yet."},
                                headers=asha_h)).json()
    assert (second["ticket_number"], second["priority"]) == ("GRV-00002", "normal")

    # Admins are notified; the list shows high priority first.
    notes = (await client.get(f"{API}/notifications", headers=admin)).json()
    assert any("URGENT Grievance GRV-00001" in n["title"] for n in notes["items"])
    listed = (await client.get(f"{API}/grievances", headers=admin)).json()
    assert [g["ticket_number"] for g in listed] == ["GRV-00001", "GRV-00002"]
    assert (await client.get(f"{API}/grievances/summary", headers=admin)).json()["high_priority_open"] == 1

    # The office replies (ticket moves to in progress); the student sees "the office", not the staff name.
    await client.post(f"{API}/grievances/{ticket['id']}/replies", json={"message": "Anti-ragging committee is meeting today."}, headers=admin)
    seen = (await client.get(f"{API}/grievances/{ticket['id']}", headers=asha_h)).json()
    assert (seen["status"], seen["replies"][0]["from_office"], seen["replies"][0]["author_name"]) == ("in_progress", True, None)
    assert (await client.get(f"{API}/grievances/{ticket['id']}", headers=admin)).json()["replies"][0]["author_name"] is not None

    resolved = (await client.put(f"{API}/grievances/{ticket['id']}/status", json={"status": "resolved", "note": "Seniors suspended."}, headers=admin)).json()
    assert (resolved["status"], resolved["resolved_at"] is not None, len(resolved["replies"])) == ("resolved", True, 2)
    # A reply from the student reopens it.
    again = (await client.post(f"{API}/grievances/{ticket['id']}/replies", json={"message": "It happened again."}, headers=asha_h)).json()
    assert (again["status"], again["resolved_at"]) == ("open", None)

    # Faculty can raise their own; nobody else can see someone's ticket.
    faculty_h = auth_headers((await login(client, "grv1-f@example.com", PASSWORD)).json()["access_token"])
    assert (await client.get(f"{API}/grievances/{ticket['id']}", headers=faculty_h)).status_code == 404
    assert (await client.get(f"{API}/grievances/mine", headers=faculty_h)).json() == []
    assert [g["ticket_number"] for g in (await client.get(f"{API}/grievances/mine", headers=asha_h)).json()] == ["GRV-00002", "GRV-00001"]
    assert (await client.post(f"{API}/grievances", json={"category": "fees", "subject": "x", "description": "too short"}, headers=asha_h)).status_code == 422


async def test_naac_report(db, client):
    admin, asha_h, asha, batch, faculty = await _college(client, "NAC1")
    await client.post(f"{API}/grievances", json={"category": "academic", "subject": "Syllabus", "description": "Syllabus not completed in time."}, headers=asha_h)

    report = await client.get(f"{API}/reports/naac", headers=admin)
    assert report.status_code == 200, report.text
    report = report.json()
    metrics = {m["label"]: m["value"] for m in report["metrics"]}
    assert (metrics["Students (current batches)"], metrics["Full-time faculty"], metrics["Student : faculty ratio"]) == (1, 1, "1.0 : 1")
    assert metrics["Grievances resolved"] == "0 of 1"
    tables = {t["key"]: t for t in report["tables"]}
    assert tables["students_programmes"]["rows"] == [["Computer Science", "B.Tech", 1, 0, 1, 0, 1]]
    assert tables["students_categories"]["rows"] == [["BC-A", 0, 1, 0, 1]]
    assert tables["faculty"]["rows"] == [["Computer Science", 1, 0, 0, 0, 1]]
    assert tables["grievances"]["rows"][0][:4] == ["Academic", 1, 0, 1]

    xlsx = await client.get(f"{API}/reports/naac.xlsx", headers=admin)
    assert xlsx.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert workbook.sheetnames[0] == "Summary" and len(workbook.sheetnames) == 1 + len(report["tables"])
    assert (await client.get(f"{API}/reports/naac", headers=asha_h)).status_code == 403
