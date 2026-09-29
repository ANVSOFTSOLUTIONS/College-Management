import io
from datetime import timedelta

from openpyxl import load_workbook

from app.db.helpers import execute, fetch_one
from app.modules.alerts.service import today_ist
from tests.factories import create_class
from tests.test_exams import _entry, _save, _setup

API = "/api/v1"


async def _mark(client, headers, class_id, day, records):
    response = await client.post(f"{API}/classes/{class_id}/attendance", json={"date": day.isoformat(), "records": records}, headers=headers)
    assert response.status_code == 200, response.text


async def test_student_attendance_report_and_excel(db, client):
    ctx = await _setup(client, "REP1")
    today = today_ist()
    s = ctx["students"]
    for back, statuses in ((1, ("present", "absent", "late")), (2, ("present", "absent", "present"))):
        records = [{"student_id": s[n]["id"], "status": st} for n, st in zip(("Asha", "Bala", "Chitra"), statuses)]
        await _mark(client, ctx["ravi"], ctx["grade"]["id"], today - timedelta(days=back), records)

    params = {"start": (today - timedelta(days=7)).isoformat(), "end": today.isoformat()}
    report = (await client.get(f"{API}/reports/student-attendance", params=params, headers=ctx["admin"])).json()
    rows = {r["full_name"]: r for r in report["rows"]}
    assert report["working_days"] == 2 and report["average_percent"] == 66.7
    assert (rows["Asha"]["percent"], rows["Bala"]["percent"], rows["Chitra"]["late"]) == (100.0, 0.0, 1)

    # The class teacher sees their class; another teacher sees nothing.
    assert len((await client.get(f"{API}/reports/student-attendance", params=params, headers=ctx["ravi"])).json()["rows"]) == 3
    assert (await client.get(f"{API}/reports/student-attendance", params=params, headers=ctx["priya"])).json()["rows"] == []
    other_class = {**params, "class_id": ctx["grade"]["id"]}
    assert (await client.get(f"{API}/reports/student-attendance", params=other_class, headers=ctx["priya"])).status_code == 404

    excel = await client.get(f"{API}/reports/student-attendance", params={**params, "format": "xlsx"}, headers=ctx["admin"])
    assert excel.headers["content-type"].startswith("application/vnd.openxmlformats")
    sheet = load_workbook(io.BytesIO(excel.content)).active
    assert sheet["D1"].value == "Student" and sheet.max_row == 4

    too_long = {"start": (today - timedelta(days=500)).isoformat(), "end": today.isoformat()}
    assert (await client.get(f"{API}/reports/student-attendance", params=too_long, headers=ctx["admin"])).status_code == 400


async def test_staff_attendance_and_fee_excel_are_admin_only(db, client):
    ctx = await _setup(client, "REP2")
    await client.post(f"{API}/staff-punch/in", headers=ctx["ravi"])
    report = (await client.get(f"{API}/reports/staff-attendance", params={"month": today_ist().isoformat()}, headers=ctx["admin"])).json()
    rows = {r["full_name"]: r for r in report["rows"]}
    assert rows["Ravi"]["present"] + rows["Ravi"]["late"] == 1 and rows["Priya"]["present"] == 0
    assert (await client.get(f"{API}/reports/staff-attendance", params={"month": today_ist().isoformat()}, headers=ctx["ravi"])).status_code == 403

    fees = await client.get(f"{API}/reports/fee-dues.xlsx", params={"only_with_dues": "false"}, headers=ctx["admin"])
    assert fees.status_code == 200 and load_workbook(io.BytesIO(fees.content)).active["A1"].value == "Class"
    assert (await client.get(f"{API}/reports/fee-dues.xlsx", headers=ctx["ravi"])).status_code == 403


async def test_exam_performance(db, client):
    ctx = await _setup(client, "REP3")
    maths, english = ctx["papers"]["Maths"], ctx["papers"]["English"]
    await _save(client, ctx["ravi"], maths, [_entry(ctx, "Asha", 45), _entry(ctx, "Bala", 10), _entry(ctx, "Chitra", 30)])
    await _save(client, ctx["priya"], english, [_entry(ctx, "Asha", 40), _entry(ctx, "Bala", 25), _entry(ctx, "Chitra", absent=True)])
    today = today_ist()
    await _mark(client, ctx["ravi"], ctx["grade"]["id"], today, [{"student_id": ctx["students"]["Bala"]["id"], "status": "absent"}])

    data = (await client.get(f"{API}/reports/exam-performance/{ctx['exam']['id']}", headers=ctx["admin"])).json()
    grade5 = data["classes"][0]
    # Asha 85%, Bala 35% (failed Maths), Chitra 30% (absent in English).
    assert (grade5["appeared"], grade5["average_percent"], grade5["pass_percent"], grade5["topper"]) == (3, 50.0, 33.3, "Asha")
    subjects = {s["subject_name"]: s for s in data["subjects"]}
    assert (subjects["Maths"]["average_percent"], subjects["Maths"]["pass_percent"]) == (56.7, 66.7)
    assert data["grades"]["A+"] == 1 and sum(data["grades"].values()) == 3
    assert [t["full_name"] for t in data["toppers"]] == ["Asha", "Bala", "Chitra"]
    attention = {a["full_name"]: a for a in data["needs_attention"]}
    assert attention["Bala"]["failed_subjects"] == ["Maths"] and attention["Bala"]["attendance_percent"] == 0.0
    assert attention["Chitra"]["absent_subjects"] == ["English"] and "Asha" not in attention

    # The class teacher may look; a subject-only teacher may not.
    assert (await client.get(f"{API}/reports/exam-performance/{ctx['exam']['id']}", headers=ctx["ravi"])).status_code == 200
    assert (await client.get(f"{API}/reports/exam-performance/{ctx['exam']['id']}", headers=ctx["priya"])).status_code == 403


async def test_attendance_follows_students_through_promotion(db, client):
    ctx = await _setup(client, "REP4")
    today = today_ist()
    s = ctx["students"]
    records = [{"student_id": s[n]["id"], "status": "present"} for n in ("Asha", "Bala", "Chitra")]
    await _mark(client, ctx["ravi"], ctx["grade"]["id"], today - timedelta(days=1), records)
    # Move Asha to a new class, as a promotion would; her earlier attendance still counts.
    grade = await fetch_one("SELECT school_id, teacher_id FROM classes WHERE id = %s", (ctx["grade"]["id"],))
    grade6_id = await create_class(school_id=grade["school_id"], teacher_id=grade["teacher_id"], name="Grade 6", academic_year="2027")
    await execute("UPDATE students SET class_id = %s WHERE id = %s", (grade6_id, s["Asha"]["id"]))
    params = {"start": (today - timedelta(days=7)).isoformat(), "end": today.isoformat(), "class_id": grade6_id}
    report = (await client.get(f"{API}/reports/student-attendance", params=params, headers=ctx["admin"])).json()
    assert [(r["full_name"], r["present"]) for r in report["rows"]] == [("Asha", 1)] and report["working_days"] == 1
