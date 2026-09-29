import io
from datetime import datetime

from openpyxl import Workbook, load_workbook

from app.modules.students.importer import COLUMNS
from tests.factories import auth_headers, create_class, create_school, create_teacher, create_user, login

API = "/api/v1"
PASSWORD = "Secret123!"
HEADERS = [f"{h}{' *' if r else ''}" for _, h, r, _ in COLUMNS]
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def _setup(client, code="IMP1"):
    school_id = await create_school(code=code)
    await create_user(school_id=school_id, email=f"{code}-admin@example.com".lower(), password=PASSWORD, role="admin")
    teacher_user = await create_user(school_id=school_id, email=f"{code}-t@example.com".lower(), password=PASSWORD, role="teacher")
    teacher_id = await create_teacher(school_id=school_id, user_id=teacher_user)
    await create_class(school_id=school_id, teacher_id=teacher_id, name="Grade 5", section="A")
    admin = auth_headers((await login(client, f"{code}-admin@example.com".lower(), PASSWORD)).json()["access_token"])
    teacher = auth_headers((await login(client, f"{code}-t@example.com".lower(), PASSWORD)).json()["access_token"])
    return admin, teacher


def _xlsx(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    for row in rows:
        sheet.append([row.get(field, "") for field, *_ in COLUMNS])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _row(**overrides):
    base = {
        "admission_number": "A-1", "full_name": "Asha Rao", "class_name": "Grade 5", "section": "A",
        "date_of_birth": "12-04-2016", "gender": "Female", "father_name": "Srinivas Rao", "father_phone": 9848022338.0,
        "mother_name": "Padma Rao", "mother_phone": "9848022339",
    }
    return {**base, **overrides}


async def _upload(client, headers, content, *, dry_run=True, skip_invalid=False, name="students.xlsx", mime=XLSX):
    return await client.post(
        f"{API}/students/import",
        params={"dry_run": str(dry_run).lower(), "skip_invalid": str(skip_invalid).lower()},
        files={"file": (name, content, mime)},
        headers=headers,
    )


async def test_template_has_columns_and_classes(db, client):
    admin, _ = await _setup(client)
    response = await client.get(f"{API}/students/import/template", headers=admin)
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert [c.value for c in workbook["Students"][1]] == HEADERS
    assert workbook["Classes"]["A3"].value == "Grade 5"


async def test_dry_run_reports_each_row_without_saving(db, client):
    admin, _ = await _setup(client)
    await client.post(f"{API}/students", json={"admission_number": "OLD-1", "full_name": "Existing", "class_id": (await client.get(f"{API}/classes", headers=admin)).json()[0]["id"]}, headers=admin)
    content = _xlsx([
        _row(),
        _row(admission_number="A-2", full_name="Bala", date_of_birth=datetime(2016, 5, 1), gender="m"),
        _row(admission_number="A-3", class_name="Grade 9", section="B"),
        _row(admission_number="A-2", full_name="Twice"),
        _row(admission_number="OLD-1"),
        _row(admission_number="A-4", date_of_birth="31/02/2016"),
        _row(admission_number="A-5", father_phone="", father_name="", mother_name="", mother_phone="", primary_contact="Uncle"),
    ])
    result = (await _upload(client, admin, content)).json()
    by_row = {r["row"]: r["errors"] for r in result["rows"]}
    assert (result["total"], result["valid"], result["invalid"], result["imported"]) == (7, 2, 5, 0)
    assert by_row[2] == [] and by_row[3] == []
    assert "not found" in by_row[4][0]
    assert "appears twice" in by_row[5][0]
    assert "already exists" in by_row[6][0]
    assert "DD-MM-YYYY" in by_row[7][0]
    assert "Primary contact" in by_row[8][0]
    assert len((await client.get(f"{API}/students", headers=admin)).json()) == 1  # nothing saved


async def test_import_saves_students_with_parents(db, client):
    admin, _ = await _setup(client)
    content = _xlsx([_row(), _row(admission_number="A-2", full_name="Bala", primary_contact="Mother")])
    result = (await _upload(client, admin, content, dry_run=False)).json()
    assert result["imported"] == 2
    students = {s["full_name"]: s for s in (await client.get(f"{API}/students", headers=admin)).json()}
    assert students["Asha Rao"]["primary_contact_phone"] == "9848022338"  # defaults to the father
    assert students["Bala"]["primary_contact_name"] == "Padma Rao"
    detail = (await client.get(f"{API}/students/{students['Asha Rao']['id']}", headers=admin)).json()
    assert (detail["date_of_birth"], detail["gender"]) == ("2016-04-12", "female")
    assert [g["relation"] for g in detail["guardians"]] == ["father", "mother"]


async def test_invalid_rows_block_import_unless_skipped(db, client):
    admin, _ = await _setup(client)
    content = _xlsx([_row(), _row(admission_number="A-2", class_name="Nope")])
    blocked = await _upload(client, admin, content, dry_run=False)
    assert blocked.status_code == 400 and blocked.json()["error"]["code"] == "rows_invalid"
    assert (await client.get(f"{API}/students", headers=admin)).json() == []
    partial = (await _upload(client, admin, content, dry_run=False, skip_invalid=True)).json()
    assert partial["imported"] == 1


async def test_csv_and_bad_files(db, client):
    admin, teacher = await _setup(client)
    csv_text = ",".join(HEADERS) + "\n" + ",".join(str(_row().get(f, "")) for f, *_ in COLUMNS).replace("9848022338.0", "9848022338") + "\n"
    ok = (await _upload(client, admin, csv_text.encode(), dry_run=False, name="s.csv", mime="text/csv")).json()
    assert ok["imported"] == 1
    assert (await _upload(client, admin, b"not excel")).status_code == 400
    missing = (await _upload(client, admin, b"Full name\nAsha\n", name="s.csv", mime="text/csv")).json()
    assert missing["error"]["code"] == "missing_columns"
    assert (await _upload(client, teacher, _xlsx([_row()]))).status_code == 403
